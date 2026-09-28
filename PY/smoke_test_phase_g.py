"""
smoke_test_phase_g.py — Gate G: the last six Group HH tools.

These are ATTACK tests, not confirmations. Every earlier gate in this project that was written
to agree with its implementation missed something an adversarial reviewer then found in
minutes, so the shape here is: assert what the tool must REFUSE, assert that a guard is not
merely present but reachable, and prefer a measurement over an assertion wherever one is cheap.

Covered:
  A  spec_builder      the prose front door writes no magnitudes and refuses to route
  B  spec_builder      the structured back door refuses a value without provenance
  C  spec_builder      editing clears acceptance and demands provenance for a magnitude
  D  plant_emitter     the emitted plant is cell-identical to the one that produced a real DLL
  E  plant_emitter     invariant I7 discriminates in BOTH directions against shipped units
  F  model_base_fork   refusals, and a no-edit fork is byte-identical
  G  unit_installer    opt-in, overlay containment, manifest-scoped uninstall
  H  hh_tools          all 24 route, none is a stub, every module loads

Run:  python PY/smoke_test_phase_g.py
"""
from __future__ import annotations

import hashlib
import json
import os
import pathlib
import shutil
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import hh_tools
import model_base_fork
import plant_emitter
import spec_builder
import spec_review
import unit_installer

REPO = pathlib.Path(__file__).resolve().parent.parent
CSTR_DIR = REPO / "Process code" if False else pathlib.Path(
    r"D:\SUMO24\Process code\Process units\10 Bioreactors\CSTR")
GOOD_CLASS = REPO / "custom_units/hh_cstr_v1/build/Plant classes/Plant.xlsx"
GOOD_INST = REPO / "custom_units/hh_cstr_v1/build/Plant instances/hh_cstr_v1.xlsx"
GOOD_UNIT = REPO / "custom_units/hh_cstr_v1/build/Process units/Custom/HH_CSTR_v1/HH_CSTR_v1.xlsx"
VENDOR_BASE = pathlib.Path(r"D:\SUMO24\Process code\Model base\Focus models\Sumo2C.xlsm")

PASS, FAIL, SKIP = [], [], []


def check(label: str, cond, detail: str = "") -> None:
    (PASS if cond else FAIL).append(label + (" -- " + detail if detail and not cond else ""))


def skip(label: str, why: str) -> None:
    SKIP.append(label + " -- " + why)


def cellmap(path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    m = {}
    for sn in wb.sheetnames:
        for row in wb[sn].iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip():
                    m[(sn, c.coordinate)] = c.value
    wb.close()
    return m


def sha(p) -> str:
    return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()


# --------------------------------------------------------------------------- #
# A — the prose front door
# --------------------------------------------------------------------------- #

def group_a() -> None:
    prose = ("A sealed tank receiving influent wastewater. Sulfide oxidation converts SH2S "
             "using oxygen SO2; the effluent leaves at the bottom and off-gas at the top. "
             "The tank is about 500 m3 and the rate constant is roughly 2.5 per day.")
    r = spec_builder.describe_to_spec(prose, "Sulfide tower")
    check("A1 prose front door returns a spec", r.get("ok") and "spec" in r)

    # THE central claim: no magnitude is ever written, even when the prose states two.
    spec = r["spec"]
    invented = [(p, n.get("value")) for p, n in spec_review.iter_values(spec)
                if spec_review._as_magnitude(n.get("value")) is not None]
    check("A2 NO magnitude invented, even though the prose states 500 and 2.5",
          not invented, "invented: %s" % invented[:4])

    # A2 alone would pass vacuously if the front door wrote a number somewhere
    # `iter_values` does not walk - there is exactly ONE value node in this spec. So scan the
    # serialised spec independently: neither number the prose states may appear anywhere in it.
    blob = json.dumps(spec)
    leaked = [tok for tok in ("500", "2.5") if tok in blob]
    check("A2b neither stated number appears ANYWHERE in the serialised spec "
          "(independent of iter_values)", not leaked, "leaked: %s" % leaked)
    check("A2c the independent scan is not vacuous - the numbers really are in the prose",
          "500" in prose and "2.5" in prose)

    check("A3 every value is the REQUIRED sentinel",
          all(n.get("value") == spec_builder.REQUIRED
              for _, n in spec_review.iter_values(spec)))
    check("A4 ports were inferred", len(spec["ports"]) >= 2)
    check("A5 species were named (structure, not quantity)",
          "SH2S" in r["inferred"]["species_mentioned"])

    # ticket 19: refuse and ask, never guess
    check("A6 reaction prose triggers a routing REFUSAL", bool(r.get("refusal")))
    check("A7 the refusal names the processes it found",
          r["refusal"] and "Sulfide oxidation" in r["refusal"]["why"])
    check("A8 the refused processes are NOT silently written into the spec",
          spec["processes"] == [])
    check("A9 routed_as stays 'A' and unconfirmed",
          spec["routing"]["routed_as"] == "A"
          and not spec["routing"]["routing_confirmed_by_user"])

    r2 = spec_builder.describe_to_spec("A simple mixing tank with an inlet and an outlet.",
                                       "Mixer")
    check("A10 non-reactive prose produces NO refusal", r2.get("ok") and not r2.get("refusal"))
    check("A11 empty prose is refused", not spec_builder.describe_to_spec("", "X").get("ok"))


# --------------------------------------------------------------------------- #
# B — the structured back door
# --------------------------------------------------------------------------- #

def group_b() -> None:
    r = spec_builder.compose_spec({"name": "Probe", "parameters": [
        {"symbol": "kX", "value": 2.5}]})
    check("B1 a value with NO provenance is REFUSED, not defaulted", not r.get("ok"))
    check("B2 the refusal names the offending parameter",
          "kX" in str(r.get("reason", "")))

    r = spec_builder.compose_spec({"name": "Probe", "parameters": [
        {"symbol": "kX", "value": 2.5, "provenance": "literature-cited",
         "source": "Henze et al. 2000, ASM2d"}]})
    check("B3 a value WITH provenance is accepted", r.get("ok"))

    r = spec_builder.compose_spec({"name": "Probe", "parameters": [{"symbol": "kX"}]})
    check("B4 an omitted value becomes REQUIRED", r.get("ok") and r["owed"] >= 1)

    check("B5 a spec with no name is refused", not spec_builder.compose_spec({}).get("ok"))
    check("B6 a port with no direction is refused",
          not spec_builder.compose_spec(
              {"name": "P", "ports": [{"symbol": "inp"}]}).get("ok"))


# --------------------------------------------------------------------------- #
# C — structured editing
# --------------------------------------------------------------------------- #

def group_c() -> None:
    spec = spec_builder.compose_spec({
        "name": "Probe", "qualified_name": "HH_Probe_v1",
        "parameters": [{"symbol": "L.Vtrain", "name": "Volume", "unit": "m3", "value": 10000,
                        "provenance": "user-measured", "source": "site survey 2026-03"}],
    })["spec"]
    paths = [u["path"] for u in spec_review.collect_unconfirmed(spec)]
    spec_review.accept(spec, "gate G", confirm=paths)
    check("C1 the fixture is accepted and may emit", spec_review.may_emit(spec)["may_emit"])

    r = spec_builder.edit_spec(spec, "parameters[0].default.value", 99999)
    check("C2 setting a magnitude with NO provenance is refused", not r.get("ok"))
    check("C3 a refused edit does not change the value",
          spec["parameters"][0]["default"]["value"] == 10000)
    check("C4 a refused edit does not clear acceptance", spec_review.may_emit(spec)["may_emit"])

    r = spec_builder.edit_spec(spec, "parameters[0].default.value", 12000,
                               provenance="user-measured", source="re-survey 2026-08")
    check("C5 a provenanced edit is applied", r.get("ok") and r["after"] == 12000)
    check("C6 the edit CLEARED acceptance", r.get("acceptance_cleared"))
    check("C7 emission now refuses until re-acceptance",
          not spec_review.may_emit(spec)["may_emit"])

    check("C8 a nonsense path is refused, not silently created",
          not spec_builder.edit_spec(spec, "nope[9].thing", 1).get("ok"))


# --------------------------------------------------------------------------- #
# D — the emitted plant, measured against one that really compiled
# --------------------------------------------------------------------------- #

def group_d(tmp: pathlib.Path) -> None:
    if not GOOD_CLASS.exists():
        skip("D1-D5", "ticket 01's reference plant is not in the tree")
        return
    root = tmp / "plant"
    r = plant_emitter.emit_test_plant("HH_CSTR_v1", root, unit_label="CSTR",
                                      instance_name="hh_cstr_v1",
                                      unit_class_path=str(GOOD_UNIT))
    check("D1 the plant emitted", r.get("emitted"), str(r.get("reason")))
    if not r.get("emitted"):
        return

    # The strongest available evidence short of another compile: the emitted class is
    # CELL-IDENTICAL to the plant that produced a verified PE32+ DLL in ticket 01.
    a, b = cellmap(GOOD_CLASS), cellmap(r["class_workbook"])
    check("D2 the emitted plant CLASS is cell-identical to the one that produced a real DLL",
          a == b, "differences: %s" % sorted(set(a) ^ set(b))[:5])
    a, b = cellmap(GOOD_INST), cellmap(r["instance_workbook"])
    check("D3 the emitted plant INSTANCE is cell-identical to that one too",
          a == b, "differences: %s" % sorted(set(a) ^ set(b))[:5])

    check("D4 the title->header invariant holds on re-read",
          r["structure_verified"]["ok"] and r["structure_verified"]["blocks_checked"] >= 3)
    check("D5 the instance symbol defaults to the class name, keeping Sumo__Plant__*",
          r["instance_symbol"] == "Plant"
          and r["variable_namespace"] == "Sumo__Plant__CSTR__*")

    for sub in plant_emitter.BUILD_SUBDIRS:
        check("D6 build root has " + sub, (root / sub).is_dir())
        break
    check("D7 all five SMT search-root subdirs exist",
          all((root / s).is_dir() for s in plant_emitter.BUILD_SUBDIRS))

    # A blank row between a title and its header is the ticket-01 defect. Prove the verifier
    # would actually catch it rather than trusting that it is wired up.
    import openpyxl
    injected = tmp / "injected.xlsx"
    shutil.copy2(r["class_workbook"], injected)
    wb = openpyxl.load_workbook(str(injected))
    ws = wb["Structure"]
    for row in ws.iter_rows():
        for c in row:
            if c.value == "Internal connection":
                ws.insert_rows(c.row + 1)
                break
    wb.save(str(injected))
    v = plant_emitter.verify_plant_workbook(injected)
    check("D8 the verifier CATCHES an injected title/header gap (the ticket-01 defect)",
          not v["ok"] and any("BLANK ROW" in f for f in v["faults"]), str(v))


# --------------------------------------------------------------------------- #
# E — invariant I7, measured against the shipped library
# --------------------------------------------------------------------------- #

def group_e(tmp: pathlib.Path) -> None:
    simple = CSTR_DIR / "Simple CSTR.xlsx"
    aer = CSTR_DIR / "CSTR with diffused aeration and calculated DO.xlsx"
    if not simple.exists() or not aer.exists():
        skip("E1-E6", "the shipped CSTR library is not reachable")
        return

    a = plant_emitter.check_reactive_precondition(simple)
    b = plant_emitter.check_reactive_precondition(aer)
    # Both directions matter. A check that only ever fails is as useless as one that only
    # ever passes; this is the negative control the earlier phases kept being asked for.
    check("E1 a bare Simple CSTR FAILS the reactive precondition", a["ok"] is False)
    check("E2 it fails on all four probe families", len(a.get("missing", [])) == 4)
    check("E3 an aeration-bearing CSTR PASSES", b["ok"] is True)
    check("E4 it passes on all four families", len(b.get("present", {})) == 4)

    r = plant_emitter.emit_test_plant("Simple CSTR", tmp / "e", reactive=True)
    check("E5 a REACTIVE plant on a bare CSTR is REFUSED", not r.get("emitted"))
    check("E6 the refusal cites invariant I7", "I7" in str(r.get("reason")))

    r = plant_emitter.emit_test_plant("CSTR with diffused aeration and calculated DO",
                                      tmp / "e2", reactive=True)
    check("E7 a REACTIVE plant on an aeration CSTR is allowed", r.get("emitted"),
          str(r.get("reason")))

    r = plant_emitter.emit_test_plant("NoSuchUnit_zzz", tmp / "e3", reactive=True)
    check("E8 reactive against an unresolvable unit is refused, not assumed fine",
          not r.get("emitted"))
    r = plant_emitter.emit_test_plant("NoSuchUnit_zzz", tmp / "e4", reactive=True,
                                      allow_unverified_reactive=True)
    check("E9 the opt-in proceeds but records that the check was SKIPPED",
          r.get("emitted") and r["precondition"]["checked"] is False)


# --------------------------------------------------------------------------- #
# F — the model-base fork
# --------------------------------------------------------------------------- #

def group_f(tmp: pathlib.Path) -> None:
    if not VENDOR_BASE.exists():
        skip("F1-F8", "the shipped model base is not reachable")
        return
    d = tmp / "fork"
    d.mkdir(parents=True, exist_ok=True)

    r = model_base_fork.fork_model_base(
        VENDOR_BASE, r"D:\SUMO24\Dir\My Process Code\Model base\My Model Category\X.xlsm")
    check("F1 refuses to write inside the install root", not r["forked"])
    r = model_base_fork.fork_model_base(VENDOR_BASE, d / "Sumo2C.xlsm")
    check("F2 refuses a name identical to a shipped model base", not r["forked"])
    check("F3 the refusal explains the shadowing mechanism",
          "shadow" in str(r.get("detail", "")).lower())
    r = model_base_fork.fork_model_base(VENDOR_BASE, d / "HH_x_v1.xlsx")
    check("F4 refuses to change .xlsm to .xlsx (it would drop the vendor macros)",
          not r["forked"])

    out = d / "HH_Sumo2C_gate_v1.xlsm"
    if out.exists():
        model_base_fork._remove(out)
    r = model_base_fork.fork_model_base(VENDOR_BASE, out, verify=True)
    check("F5 a no-edit fork succeeds", r["forked"], str(r.get("reason")))
    if r["forked"]:
        check("F6 a no-edit fork is BYTE-IDENTICAL to the vendor original",
              sha(VENDOR_BASE) == sha(out))
        check("F7 no OOXML part was lost (this is what an openpyxl round trip destroys)",
              r["parts_after_copy"]["n_lost"] == 0
              and r["parts_after_copy"]["n_vendor_parts"] > 60)
        check("F8 the fork is WRITABLE (a copy inherits the vendor read-only bit)",
              r["writable"] and os.access(str(out), os.W_OK))
        check("F9 rung 2.5 was run on the fork and passed additively",
              r["additive_only_gate"]["ok"] and r["diff_summary"]["n_modified"] == 0)

    r = model_base_fork.fork_model_base(VENDOR_BASE, out)
    check("F10 refuses to overwrite an existing fork", not r["forked"])

    # The COM path is exercised only when Excel is present; say so rather than pretending.
    try:
        import win32com.client  # noqa: F401
        have_com = True
    except ImportError:
        have_com = False
    if not have_com:
        skip("F11-F12", "Excel COM (pywin32) unavailable; the insertion path was not exercised")
        return
    bad = d / "HH_Sumo2C_gate_v2.xlsm"
    if bad.exists():
        model_base_fork._remove(bad)
    r = model_base_fork.fork_model_base(VENDOR_BASE, bad, additions={
        "Parameters": [{"Symbol": "kGate", "NoSuchColumn": 1}]})
    check("F11 a bad column name is refused AND the partial fork is removed",
          not r["forked"] and r.get("partial_fork_removed") and not bad.exists())


# --------------------------------------------------------------------------- #
# G — install
# --------------------------------------------------------------------------- #

def group_g(tmp: pathlib.Path) -> None:
    real_root = unit_installer.OVERLAY_ROOT
    real_units = unit_installer.OVERLAY_UNITS
    group_g._overlay_before = sorted(str(p) for p in real_units.rglob("*.xlsx")) \
        if real_units.exists() else []

    r = unit_installer.install_process_unit("anything")
    check("G1 install without opt_in is refused", not r["installed"])

    try:
        unit_installer._assert_overlay(pathlib.Path(r"D:\SUMO24\Process code\Process units\X"))
        check("G2 the overlay guard blocks a vendor-tree path", False, "it did NOT raise")
    except unit_installer.InstallPathError:
        check("G2 the overlay guard blocks a vendor-tree path", True)

    if not GOOD_UNIT.exists():
        skip("G3-G10", "no reference unit workbook to stage")
        return

    # Redirect the overlay so the gate never writes into D:\SUMO24.
    overlay = tmp / "overlay"
    unit_installer.OVERLAY_ROOT = overlay
    unit_installer.OVERLAY_UNITS = overlay / "Process Units"
    unit_installer.OVERLAY_MODEL_BASE = overlay / "Model base"
    overlay.mkdir(parents=True, exist_ok=True)
    try:
        slug = tmp / "slug" / "HH_Gate_v1"
        udir = slug / "build" / "Process units" / "Custom" / "HH_Gate_v1"
        udir.mkdir(parents=True, exist_ok=True)
        (slug / "artifacts").mkdir(parents=True, exist_ok=True)
        shutil.copy2(GOOD_UNIT, udir / "HH_Gate_v1.xlsx")
        (udir / "HH_Gate_v1.emf").write_bytes(b"\x01\x00\x00\x00" + b"\x00" * 64)
        (udir.parent / "Custom Group Info.xlsx").write_bytes(b"PK\x03\x04")
        (slug / "artifacts" / "g.dll").write_bytes(b"MZ" + b"\x00" * 64)

        spec = spec_builder.compose_spec({
            "name": "Gate probe", "qualified_name": "HH_Gate_v1",
            "parameters": [{"symbol": "L.Vtrain", "name": "Volume", "unit": "m3",
                            "value": 1000, "provenance": "user-measured",
                            "source": "gate fixture"}]})["spec"]

        (slug / "unit_spec.json").write_text(json.dumps(spec), encoding="utf-8")
        pre = unit_installer.preflight(slug)
        check("G3 an UNACCEPTED spec blocks install at preflight", not pre["ok"])

        paths = [u["path"] for u in spec_review.collect_unconfirmed(spec)]
        spec_review.accept(spec, "gate G", confirm=paths)
        (slug / "unit_spec.json").write_text(json.dumps(spec), encoding="utf-8")
        pre = unit_installer.preflight(slug)
        check("G4 an accepted spec passes preflight", pre["ok"], str(pre["reason"]))

        d = unit_installer.install_process_unit(slug, opt_in=True, dry_run=True)
        check("G5 dry run writes nothing and lists what it would write",
              not d["installed"] and len(d["would_write"]) == 3
              and not (unit_installer.OVERLAY_UNITS).exists())

        i = unit_installer.install_process_unit(slug, opt_in=True)
        check("G6 install writes the unit, icon and group info", i["installed"]
              and i["n_files"] == 3, str(i.get("reason")))
        check("G6b install maps onto category/<family>/ with the Group Info BESIDE the "
              "workbook (measured 2026-09-07)",
              i["installed"]
              and pathlib.Path(i["destination"]).name == "Custom"
              and (pathlib.Path(i["destination"]) / "Custom Group Info.xlsx").exists(),
              str(i.get("destination")))
        check("G7 install writes a manifest",
              (slug / unit_installer.MANIFEST_NAME).exists())
        check("G8 the result warns that the GUI re-scans at STARTUP ONLY",
              "STARTUP ONLY" in i.get("IMPORTANT", ""))

        i2 = unit_installer.install_process_unit(slug, opt_in=True)
        check("G9 a second install of the same version is REFUSED", not i2["installed"])

        # Uninstall must remove what it installed and KEEP what someone else changed.
        touched = pathlib.Path(i["files"][0])
        touched.write_bytes(touched.read_bytes() + b"EDITED-BY-SOMEONE-ELSE")
        u = unit_installer.uninstall_process_unit(slug, opt_in=True)
        check("G10 uninstall removes the untouched files", len(u["removed"]) == 2)
        check("G11 uninstall KEEPS a file changed since install", len(u["kept"]) == 1
              and touched.exists())
        check("G12 uninstall without opt_in is refused",
              not unit_installer.uninstall_process_unit(slug)["uninstalled"])
    finally:
        unit_installer.OVERLAY_ROOT = real_root
        unit_installer.OVERLAY_UNITS = real_units
        unit_installer.OVERLAY_MODEL_BASE = real_root / "Model base"

    # The gate must leave the real install exactly as it found it. [fixed 2026-09-07]
    # The overlay legitimately holds installed units; assert nothing was ADDED, not that
    # it is empty.
    live_after = sorted(str(p) for p in real_units.rglob("*.xlsx")) \
        if real_units.exists() else []
    check("G13 the real D:\\SUMO24 overlay was never written to",
          live_after == sorted(str(p) for p in getattr(group_g, "_overlay_before", live_after)),
          "delta: %s" % sorted(set(live_after) -
                               set(getattr(group_g, "_overlay_before", []))))


# --------------------------------------------------------------------------- #
# H — the MCP surface
# --------------------------------------------------------------------------- #

def group_h() -> None:
    check("H1 24 tools are defined", len(hh_tools.TOOL_DEFS) == 24,
          "got %d" % len(hh_tools.TOOL_DEFS))
    check("H2 no duplicate tool names",
          len(set(hh_tools.TOOL_NAMES)) == len(hh_tools.TOOL_NAMES))
    check("H3 every rung module imported", not hh_tools._IMPORT_ERRORS,
          str(hh_tools._IMPORT_ERRORS))

    unrouted, stubs = [], []
    for n in hh_tools.TOOL_NAMES:
        r = hh_tools.dispatch(n, {})
        if str(r.get("error", "")).startswith("unknown Group HH tool"):
            unrouted.append(n)
        if r.get("not_implemented"):
            stubs.append(n)
    check("H4 every tool routes through dispatch", not unrouted, str(unrouted))
    check("H5 NO tool is still a not-implemented stub", not stubs, str(stubs))

    # dispatch must never raise across the MCP boundary (tickets 17/19)
    raised = []
    for n in hh_tools.TOOL_NAMES:
        try:
            out = hh_tools.dispatch(n, {"spec_path": "/nope", "vendor": "/nope",
                                        "out_path": "/nope", "xml": "/nope"})
            if not isinstance(out, dict):
                raised.append(n + " (non-dict)")
        except Exception as e:                               # pragma: no cover
            raised.append("%s (%s)" % (n, type(e).__name__))
    check("H6 no tool raises across the MCP boundary on bad input", not raised, str(raised))

    check("H7 the new modules appear in the environment report",
          all(k in hh_tools._check_environment()["rung_modules"]
              for k in ("spec_builder", "plant_emitter", "model_base_fork", "unit_installer")))


def main() -> int:
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="gate_g_"))
    try:
        group_a()
        group_b()
        group_c()
        group_d(tmp)
        group_e(tmp)
        group_f(tmp)
        group_g(tmp)
        group_h()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    for line in PASS:
        print("  ok   " + line)
    for line in SKIP:
        print("  skip " + line)
    for line in FAIL:
        print("  FAIL " + line)
    total = len(PASS) + len(FAIL)
    print("\nGate G: %d/%d passed, %d skipped" % (len(PASS), total, len(SKIP)))
    return 0 if not FAIL else 1


if __name__ == "__main__":
    sys.exit(main())
