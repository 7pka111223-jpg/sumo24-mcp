"""
model_base_diff.py — Group HH, ticket 12 Phase B2: rung 2.5, the model-base semantic diff.

Type B forks a whole model base (ticket 04) and edits it ADDITIVELY. This rung is what makes
"additively" checkable rather than aspirational, and it is the **load-bearing condition of
ticket 18's calibrated-plant policy**: adding a Type B unit to a calibrated plant is permissible
ONLY behind a passing additive-only diff.

WHAT IT COMPARES, and why this way:

  * **By SYMBOL NAME, never by row position.** A fork inserts rows; every position shifts.
    Positional diffing would report the entire remainder of a sheet as modified. This is the
    same lesson that produced header-name lookup in rung 2 and name-keyed flux comparison in
    rung 5 - three rungs, one recurring defect class.

  * **Only the value-bearing sheets.** `Parameters`, `other parameters`, `Components`,
    `Calculated variables`, `systemcode_constants` carry symbol->value rows; `Model` carries the
    Gujer matrix keyed by process id. Help/Concepts/Security are prose and are ignored - a
    documentation edit is not a model change.

  * **Three verdicts, and only one of them is allowed:** `added` (fine), `modified` (FAILS the
    gate), `removed` (FAILS the gate). A modified vendor coefficient is exactly the silent
    result-corrupting change Type B was accepted "with teeth" to prevent (map Q14).

  * **Reading uses openpyxl; WRITING a fork does not.** [MEASURED, ticket 06] An openpyxl
    round-trip of a model base drops ~124 symbol-map entries through sharedStrings loss. Reading
    is safe; the fork writer is Excel COM (ticket 11), never openpyxl.
"""
from __future__ import annotations

import pathlib
from typing import Any

# Sheets that carry symbol -> value rows worth diffing. Prose sheets are deliberately excluded.
VALUE_SHEETS = ("Parameters", "other parameters", "Components", "Calculated variables",
                "systemcode_constants", "PU parameters")
MATRIX_SHEET = "Model"


def _cells(row) -> list[str]:
    return [("" if c is None else str(c).strip()) for c in row]


def read_value_sheet(ws) -> dict[str, dict[str, Any]]:
    """`{symbol: {column_header: value}}` for a value-bearing sheet.

    Blocks repeat down the sheet: a title row, a header row starting with `Symbol`, then data
    rows until a blank. Headers are re-read per block because they vary (the `Default 1 / Default
    i(2 to n)` array variants), so a single global header assumption would mis-key later blocks.
    """
    out: dict[str, dict[str, Any]] = {}
    header: list[str] | None = None
    sym_col = 1
    for row in ws.iter_rows(values_only=True):
        cells = _cells(row)
        if not any(cells):
            header = None
            continue
        if "Symbol" in cells:
            sym_col = cells.index("Symbol")
            header = cells
            continue
        if header is None:
            continue
        if sym_col >= len(cells):
            continue
        sym = cells[sym_col]
        if not sym or sym == "Symbol":
            continue
        rec = {}
        for i, h in enumerate(header):
            if h and h != "Symbol" and i < len(cells):
                rec[h] = cells[i]
        out[sym] = rec
    return out


def read_gujer(ws) -> dict[str, dict[str, str]]:
    """`{process_id: {component: coefficient}}` from the SYMBOLIC Gujer block.

    Keyed by the `r<j>` process id in the `Symbol` column, with components taken from the header
    row - so an inserted process shifts nothing.
    """
    out: dict[str, dict[str, str]] = {}
    header: list[str] | None = None
    sym_col = 2
    for row in ws.iter_rows(values_only=True):
        cells = _cells(row)
        if not any(cells):
            continue
        if "Symbol" in cells and "Name" in cells:
            header = cells
            sym_col = cells.index("Symbol")
            continue
        if header is None or sym_col >= len(cells):
            continue
        pid = cells[sym_col]
        if not pid.startswith("r") or not pid[1:].split(",")[0].isdigit():
            continue
        rec = {}
        for i, h in enumerate(header):
            if h and h not in ("j", "Symbol", "Name") and i < len(cells) and cells[i]:
                rec[h] = cells[i]
        if pid in out:            # the evaluated second block repeats ids; keep the first
            continue
        out[pid] = rec
    return out


def read_model_base(path: str | pathlib.Path) -> dict[str, dict]:
    """Read the value sheets + the Gujer matrix. READ-ONLY; openpyxl is safe for reading."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    data: dict[str, dict] = {}
    for sn in VALUE_SHEETS:
        if sn in wb.sheetnames:
            data[sn] = read_value_sheet(wb[sn])
    if MATRIX_SHEET in wb.sheetnames:
        data[MATRIX_SHEET] = read_gujer(wb[MATRIX_SHEET])
    wb.close()
    return data


def diff_model_bases(vendor: str | pathlib.Path, fork: str | pathlib.Path) -> dict:
    """Semantic diff, keyed by symbol. Returns added / modified / removed per sheet."""
    v, f = read_model_base(vendor), read_model_base(fork)
    added, modified, removed = [], [], []
    for sheet in sorted(set(v) | set(f)):
        vs, fs = v.get(sheet, {}), f.get(sheet, {})
        for sym in sorted(set(fs) - set(vs)):
            added.append({"sheet": sheet, "symbol": sym})
        for sym in sorted(set(vs) - set(fs)):
            removed.append({"sheet": sheet, "symbol": sym})
        for sym in sorted(set(vs) & set(fs)):
            a, b = vs[sym], fs[sym]
            for key in sorted(set(a) | set(b)):
                av, bv = a.get(key, ""), b.get(key, "")
                if av != bv:
                    modified.append({"sheet": sheet, "symbol": sym, "field": key,
                                     "vendor": av, "fork": bv})
    return {"added": added, "modified": modified, "removed": removed,
            "n_added": len(added), "n_modified": len(modified), "n_removed": len(removed),
            "sheets_compared": sorted(set(v) | set(f))}


def additive_only_gate(diff: dict) -> dict:
    """THE GATE. Passes only when the fork is purely additive.

    Ticket 18's calibrated-plant policy rests entirely on this: a `modified` or `removed` entry
    means the fork changed something a calibrated plant already depends on, and the change would
    propagate silently into results that were previously validated.
    """
    bad = diff.get("n_modified", 0) + diff.get("n_removed", 0)
    if bad:
        return {"ok": False, "gate": "additive-only",
                "reason": f"fork is NOT additive: {diff['n_modified']} modified, "
                          f"{diff['n_removed']} removed",
                "detail": ["A calibrated plant's results depend on the vendor values; a modified "
                           "or removed entry changes them silently.",
                           "Only ADDED symbols are permissible for a Type B fork."],
                "modified": diff.get("modified", [])[:8],
                "removed": diff.get("removed", [])[:8],
                "n_added": diff.get("n_added", 0)}
    return {"ok": True, "gate": "additive-only",
            "reason": f"purely additive: {diff.get('n_added', 0)} symbol(s) added, "
                      f"nothing modified or removed",
            "n_added": diff.get("n_added", 0), "added": diff.get("added", [])[:8]}


# --------------------------------------------------------------------------- #
# Q14 teeth 3 and 4 — the Type B safeguards that were decided but never built
# --------------------------------------------------------------------------- #

def stamp_verification_state(spec: dict, continuity_ok: bool,
                             kinetics_validated: bool = False) -> dict:
    """Q14 tooth 3: stamp what has and has NOT been verified.

    Stoichiometry can be proven by rung 2. Kinetics cannot be proven by any rung here - only by
    data. A Type B unit whose rates were never validated must SAY SO on the artefact, so a plant
    carrying it cannot be mistaken for a validated one.
    """
    state = ("stoichiometry-verified, kinetically-unverified" if continuity_ok and not kinetics_validated
             else "stoichiometry-verified, kinetically-validated" if continuity_ok and kinetics_validated
             else "stoichiometry-UNVERIFIED")
    # The schema ALREADY defines `provenance.machine_stamp` for exactly this
    # ("stoichiometry-verified, kinetically-unverified until the degenerate test runs").
    # An earlier version invented `verification_state`, which additionalProperties:false
    # would have rejected - found by building the spec tools against the real schema.
    spec.setdefault("provenance", {})["machine_stamp"] = state
    return {"stamped": state, "continuity_ok": continuity_ok,
            "kinetics_validated": kinetics_validated}


def type_b_admission(spec: dict, gate_result: dict, continuity_ok: bool) -> dict:
    """Q14 tooth 4: a Type B unit that fails its gate is DISABLED, not silently included.

    `gate_result` is passed IN, never recomputed here (ticket 18): a tool that re-runs its own
    gate can be made to skip it. The caller must supply the evidence.
    """
    routed_as = (spec.get("routing") or {}).get("routed_as")
    if routed_as != "B":
        return {"admitted": True, "reason": "not a Type B unit; rung 2.5 does not apply"}
    if not isinstance(gate_result, dict) or gate_result.get("gate") != "additive-only":
        return {"admitted": False, "disabled": True,
                "reason": "no additive-only gate result supplied",
                "detail": ["Ticket 18: the gate result is an INPUT, so it cannot be skipped. "
                           "Run additive_only_gate() and pass its result."]}
    if not gate_result.get("ok"):
        return {"admitted": False, "disabled": True,
                "reason": "the fork failed the additive-only gate",
                "detail": [gate_result.get("reason", "")]}
    if not continuity_ok:
        return {"admitted": False, "disabled": True,
                "reason": "continuity (rung 2) did not pass for this unit"}
    stamp = stamp_verification_state(spec, continuity_ok=True, kinetics_validated=False)
    return {"admitted": True, "reason": "additive fork, continuity clean",
            "verification_state": stamp["stamped"],
            "caveat": "kinetics remain unvalidated - the stamp says so on the artefact"}


if __name__ == "__main__":                                   # pragma: no cover
    print(__doc__)
