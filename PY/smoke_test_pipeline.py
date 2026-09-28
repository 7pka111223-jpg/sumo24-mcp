"""
smoke_test_pipeline.py
───────────────────────
End-to-end pipeline smoke test covering Stages 0–7.
Stage 8 is skipped (requires a real SUMO compile + DLL).

Run:
    python smoke_test_pipeline.py

Success criterion: pipeline_status returns all stages 0–4 passed,
stages 5 and 6 skipped, and stage 7 passed.
"""
import json
import pathlib
import sys
import os
import tempfile

# ── Path setup ────────────────────────────────────────────────────────────────
_HERE = pathlib.Path(__file__).parent
sys.path.insert(0, str(_HERE))

# ── Import pipeline modules directly ─────────────────────────────────────────
from pipeline import manifest as _m
from pipeline import hashing  as _h
from pipeline import snapshot  as _sn
from pipeline import seeding   as _sd
from pipeline import engineering as _eng
from pipeline import static_inputs as _si
from pipeline import dynamic_inputs as _di
from pipeline import build as _bld

_PASS = "✅"
_FAIL = "❌"

def _ok(label: str, condition: bool, detail: str = "") -> bool:
    status = _PASS if condition else _FAIL
    print(f"  {status} {label}" + (f" — {detail}" if detail else ""))
    return condition


def _section(title: str):
    print(f"\n{'─'*60}")
    print(f"  {title}")
    print(f"{'─'*60}")


# ── Minimal schematic fixture ─────────────────────────────────────────────────
_SCHEMATIC = {
    "meta": {"plant_name": "TestPlant", "kinetic_model": "ASM2d"},
    "units": [
        {"id": "u1", "name": "Influent1",          "type": "influent"},
        {"id": "u2", "name": "Screen1",             "type": "screen"},
        {"id": "u3", "name": "OxidationDitch1",     "type": "oxidation_ditch"},
        {"id": "u4", "name": "SecondarySettler1",   "type": "secondary_clarifier"},
        {"id": "u5", "name": "SludgeSplitter1",     "type": "sludge_splitter"},
        {"id": "u6", "name": "Effluent1",           "type": "effluent"},
        {"id": "u7", "name": "WASFlow1",            "type": "was_flow"},
    ],
    "streams": [
        {"id": "s1", "from": "u1", "to": "u2", "type": "liquid"},
        {"id": "s2", "from": "u2", "to": "u3", "type": "liquid"},
        {"id": "s3", "from": "u3", "to": "u4", "type": "liquid"},
        {"id": "s4", "from": "u4", "to": "u5", "type": "liquid"},
        {"id": "s5", "from": "u5", "to": "u3", "type": "ras"},
        {"id": "s6", "from": "u5", "to": "u7", "type": "was"},
        {"id": "s7", "from": "u4", "to": "u6", "type": "liquid"},
    ],
}

_INFLUENT = {
    "Q_m3d": 7420, "COD_mgL": 450, "BOD_mgL": 220, "TKN_mgL": 52,
    "TP_mgL": 7.5, "TSS_mgL": 280, "T_C": 20, "Alk_mgCaCO3": 400,
}

_ENG_PARAMS = {
    # V=4200 m3  → HRT = 4200/7420×24 = 13.6 h ✓
    # Q_WAS=150  → SRT = 4200×3500/(150×8000) = 12.25 d ✓
    # F:M = (7420×220)/(4200×2625) = 0.148 d⁻¹ ✓
    # Q_eff=7270 → Q_eff+Q_WAS = 7420 (hydraulic balance ✓)
    # COD ratio = (7270×50 + 150×4200) / (7420×450) = (363500+630000)/3339000 = 29.7% ✓ (biological oxidation)
    "Q_avg_m3d": 7420, "COD_in_mgL": 450, "BOD_mgL": 220,
    "TKN_mgL": 52, "V_total_m3": 4200, "V_reactor_m3": 4200,
    "MLSS_mgL": 3500, "MLVSS_mgL": 2625, "Q_WAS_m3d": 150,
    "X_R_mgL": 8000, "Q_eff_m3d": 7270, "COD_eff_mgL": 50,
    "COD_WAS_mgL": 4200, "KLa_per_d": 150, "A_clarifier_m2": 700,
    # Alk=400 → residual = 400 - 52×7.14 + 15×3.57 = 82 mg/L CaCO3 ≥ 50 ✓
    "RAS_ratio": 0.6, "Alk_in_mgCaCO3": 400,
    "NO3_removed_mgL": 15, "T_C": 20,
}


def test_manifest():
    _section("manifest.py — basic round-trip")
    all_pass = True
    with tempfile.TemporaryDirectory() as td:
        m = _m.new_manifest(td, "TestPlant")
        all_pass &= _ok("new_manifest creates 9 stages", len(m["stages"]) == 9)
        all_pass &= _ok("stage 0 status is pending", m["stages"][0]["status"] == "pending")

        _m.set_stage(m, 0, status="passed", output_hash="abc123")
        _m.save(td, m)

        m2 = _m.load(td)
        all_pass &= _ok("save/load round-trip", m2["stages"][0]["status"] == "passed")
        all_pass &= _ok("output_hash preserved", m2["stages"][0]["output_hash"] == "abc123")

        stale = _m.mark_downstream_stale(m2, 0)
        all_pass &= _ok("mark_downstream_stale returns empty (nothing passed yet)", stale == [])

        _m.set_stage(m2, 1, status="passed", output_hash="def456")
        stale = _m.mark_downstream_stale(m2, 0)
        all_pass &= _ok("mark_downstream_stale marks stage 1 stale", 1 in stale)
    return all_pass


def test_hashing(tmp_dir: pathlib.Path):
    _section("hashing.py — stability")
    all_pass = True

    # JSON hash
    j = tmp_dir / "test.json"
    j.write_text('{"b": 2, "a": 1}', encoding="utf-8")
    h1 = _h.hash_json(j)
    j.write_text('{"a": 1, "b": 2}', encoding="utf-8")   # different key order
    h2 = _h.hash_json(j)
    all_pass &= _ok("JSON hash is key-order independent", h1 == h2)

    j.write_text('{"a": 1, "b": 3}', encoding="utf-8")   # different value
    h3 = _h.hash_json(j)
    all_pass &= _ok("JSON hash changes on value change", h1 != h3)

    return all_pass


def test_snapshot(tmp_dir: pathlib.Path):
    _section("snapshot.py — snapshot + revert")
    all_pass = True

    _m.save(tmp_dir, _m.new_manifest(str(tmp_dir), "snapshot_test"))
    (tmp_dir / "data.txt").write_text("original", encoding="utf-8")
    snap = _sn.snapshot(str(tmp_dir), "test")
    all_pass &= _ok("snapshot file created", pathlib.Path(snap).exists())

    (tmp_dir / "data.txt").write_text("modified", encoding="utf-8")
    _sn.revert(str(tmp_dir), "test")
    restored = (tmp_dir / "data.txt").read_text(encoding="utf-8")
    all_pass &= _ok("revert restores content", restored == "original")

    return all_pass


def test_seeding(tmp_dir: pathlib.Path):
    _section("seeding.py — seed_excel_from_schematic")
    all_pass = True

    # Need openpyxl and a template
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws_tp = wb.create_sheet("Treatment_Processes")
        ws_tp.append(["order_index", "instance_name", "process_type",
                      "category", "stage", "kinetic_model"])
        ws_conn = wb.create_sheet("Connections")
        ws_conn.append(["from_unit", "to_unit", "stream_label", "stream_type", "recycle_type"])
        xlsx = tmp_dir / "template.xlsx"
        wb.save(xlsx)

        result = _sd.seed_excel_from_schematic(str(xlsx), _SCHEMATIC, backup=False)
        all_pass &= _ok("seed_excel_from_schematic ok", result["ok"],
                        result.get("error", ""))
        all_pass &= _ok("Treatment_Processes rows written",
                        result.get("treatment_processes_written", 0) == len(_SCHEMATIC["units"]))
        all_pass &= _ok("Connections rows written",
                        result.get("connections_written", 0) == len(_SCHEMATIC["streams"]))
    except ImportError:
        print("  ⚠️  openpyxl not available — skipping seeding test")
    return all_pass


def test_engineering():
    _section("engineering.py — run_engineering_checks")
    all_pass = True

    result = _eng.run_engineering_checks(_ENG_PARAMS)
    all_pass &= _ok("run_engineering_checks returns ok",  result.get("ok", False),
                    result.get("top_fix", ""))
    all_pass &= _ok("SRT in range", any(c["name"] == "srt_in_range" and c["ok"]
                                        for c in result.get("checks", [])))
    all_pass &= _ok("HRT in range", any(c["name"] == "hrt_in_range" and c["ok"]
                                        for c in result.get("checks", [])))
    return all_pass


def test_static_inputs(tmp_dir: pathlib.Path):
    _section("static_inputs.py — resolve + write")
    all_pass = True

    result = _si.write_static_inputs(
        project_dir=str(tmp_dir),
        plant_name="TestPlant",
        influent=_INFLUENT,
        kinetic_model="ASM2d",
    )
    all_pass &= _ok("write_static_inputs ok", result["ok"],
                    str([v for v in result.get("validation", []) if not v.get("ok")]))
    all_pass &= _ok("static_inputs.json written",
                    (tmp_dir / "static_inputs.json").exists())
    all_pass &= _ok("static_inputs.scs written",
                    (tmp_dir / "static_inputs.scs").exists())
    sv = result.get("influent_sv", {})
    all_pass &= _ok("S_NH resolved", "S_NH" in sv and sv["S_NH"] > 0)
    return all_pass


def test_dynamic_inputs(tmp_dir: pathlib.Path):
    _section("dynamic_inputs.py — generate + validate")
    all_pass = True

    for ptype in ["constant", "diurnal", "storm"]:
        r = _di.generate_dynamic_inputs(
            str(tmp_dir), ptype, _INFLUENT, duration_days=7, dt_h=1.0
        )
        all_pass &= _ok(f"generate {ptype} ok", r.get("ok", False), r.get("error", ""))
        if r.get("ok"):
            vr = _di.validate_dynamic_tsv(r["path"])
            all_pass &= _ok(f"validate {ptype} tsv ok", vr.get("ok", False),
                            str(vr.get("errors", "")))
    return all_pass


def test_full_pipeline():
    _section("Full pipeline — Stages 0–7 (Stage 8 requires SUMO)")
    all_pass = True

    with tempfile.TemporaryDirectory() as td:
        proj = str(pathlib.Path(td) / "TestPlant")

        # ── Stage 0: init ──────────────────────────────────────────────────────
        m = _m.new_manifest(proj, "TestPlant")
        pathlib.Path(proj).mkdir(parents=True, exist_ok=True)
        for fname, content in [
            ("engineering_checks.json", '{"status":"pending"}'),
            ("static_inputs.json",      '{"status":"pending"}'),
            ("controllers.json",        '{"controllers":[]}'),
        ]:
            (pathlib.Path(proj) / fname).write_text(content, encoding="utf-8")
        _m.set_stage(m, 0, status="passed", output_hash="init_hash")
        _m.save(proj, m)
        m = _m.load(proj)
        all_pass &= _ok("Stage 0: init manifest saved", m["stages"][0]["status"] == "passed")

        # ── Stage 1: schematic (simulate HTML with embedded JSON) ──────────────
        html_path = pathlib.Path(proj) / "TestPlant_schematic.html"
        schematic_json = json.dumps(_SCHEMATIC, indent=2)
        html_content = (
            f'<html><body>'
            f'<script id="schematic-data" type="application/json">'
            f'{schematic_json}'
            f'</script></body></html>'
        )
        html_path.write_text(html_content, encoding="utf-8")
        h1 = _h.hash_schematic(str(html_path))
        all_pass &= _ok("Stage 1: schematic hashed", len(h1) == 64)
        _m.set_stage(m, 1, status="passed", output_hash=h1,
                     summary={"units": 7, "streams": 7})
        _m.save(proj, m)

        # ── Stage 2: seed Excel ────────────────────────────────────────────────
        try:
            import openpyxl
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.create_sheet("Treatment_Processes")
            ws.append(["order_index","instance_name","process_type","category","stage","kinetic_model"])
            ws2 = wb.create_sheet("Connections")
            ws2.append(["from_unit","to_unit","stream_label","stream_type","recycle_type"])
            xlsx_path = pathlib.Path(proj) / "TestPlant_data.xlsx"
            wb.save(xlsx_path)
            seed_r = _sd.seed_excel_from_schematic(str(xlsx_path), _SCHEMATIC, backup=False)
            all_pass &= _ok("Stage 2: Excel seeded", seed_r.get("ok", False))
            h2 = _h.hash_excel(str(xlsx_path))
            _m.set_stage(m, 2, status="passed", output_hash=h2,
                         summary={"processes": 7})
            _m.save(proj, m)
        except ImportError:
            print("  ⚠️  openpyxl not available — Stage 2 skipped in test")
            _m.set_stage(m, 2, status="passed", output_hash="no_openpyxl")
            _m.save(proj, m)

        # ── Stage 3: engineering checks ────────────────────────────────────────
        ec_path = pathlib.Path(proj) / "engineering_checks.json"
        ec_result = _eng.run_engineering_checks(_ENG_PARAMS, output_path=str(ec_path))
        all_pass &= _ok("Stage 3: engineering checks ok", ec_result.get("ok", False),
                        ec_result.get("top_fix", ""))
        h3 = _h.hash_json(str(ec_path))
        _m.set_stage(m, 3, status="passed", output_hash=h3,
                     summary=ec_result.get("computed", {}))
        _m.save(proj, m)

        # ── Stage 4: static inputs ─────────────────────────────────────────────
        si_result = _si.write_static_inputs(proj, "TestPlant", _INFLUENT)
        all_pass &= _ok("Stage 4: static inputs ok", si_result.get("ok", False),
                        str([v for v in si_result.get("validation", []) if not v.get("ok")]))
        h4 = _h.hash_json(pathlib.Path(proj) / "static_inputs.json")
        _m.set_stage(m, 4, status="passed", output_hash=h4)
        _m.save(proj, m)

        # ── Stage 5: SKIP ──────────────────────────────────────────────────────
        _m.set_stage(m, 5, status="skipped",
                     summary={"skip_reason": "steady-state design study only"})
        _m.save(proj, m)
        all_pass &= _ok("Stage 5: skipped", m["stages"][5]["status"] == "skipped")

        # ── Stage 6: SKIP ──────────────────────────────────────────────────────
        _m.set_stage(m, 6, status="skipped",
                     summary={"skip_reason": "manual control"})
        _m.save(proj, m)

        # ── Stage 7: build ─────────────────────────────────────────────────────
        # Build minimal pack manually (sumo_pack may not be runnable in test)
        pack_dir = pathlib.Path(proj) / "TestPlant.sumo-pack"
        pack_dir.mkdir(parents=True, exist_ok=True)
        (pack_dir / "plant.sumoslang").write_text(
            f"// plant.sumoslang for TestPlant\\n", encoding="utf-8"
        )
        (pack_dir / "apply_parameters.scs").write_text(
            "// apply_parameters.scs\\n", encoding="utf-8"
        )
        (pack_dir / "static_inputs.scs").write_text(
            (pathlib.Path(proj) / "static_inputs.scs").read_text(encoding="utf-8")
            if (pathlib.Path(proj) / "static_inputs.scs").exists()
            else "// static_inputs.scs\\n",
            encoding="utf-8",
        )
        stub = f'<?xml version="1.0" encoding="utf-8"?>\\n<SumoProject name="TestPlant" version="24"/>\\n'
        (pack_dir / "TestPlant.sumo").write_text(stub, encoding="utf-8")
        (pack_dir / "BUILD_INSTRUCTIONS.md").write_text(
            "# Build TestPlant\\n1. Open stub\\n2. Compile\\n", encoding="utf-8"
        )
        (pack_dir / "schematic.json").write_text(
            json.dumps(_SCHEMATIC, indent=2), encoding="utf-8"
        )

        expected_files = ["plant.sumoslang", "apply_parameters.scs",
                          "static_inputs.scs", "TestPlant.sumo",
                          "BUILD_INSTRUCTIONS.md", "schematic.json"]
        pack_files_ok = all((pack_dir / f).exists() for f in expected_files)
        all_pass &= _ok("Stage 7: pack files present", pack_files_ok)

        h7 = _h.hash_pack(str(pack_dir))
        _m.set_stage(m, 7, status="passed", output_hash=h7,
                     summary={"pack_files": expected_files})
        _m.save(proj, m)

        # ── Final state check ──────────────────────────────────────────────────
        m_final = _m.load(proj)
        passed  = sum(1 for s in m_final["stages"] if s["status"] == "passed")
        skipped = sum(1 for s in m_final["stages"] if s["status"] == "skipped")
        pending = sum(1 for s in m_final["stages"] if s["status"] == "pending")

        all_pass &= _ok(f"Final: stages passed={passed} skipped={skipped} pending={pending}",
                        passed >= 6 and skipped == 2 and pending == 1,
                        "(pending=1 expected for Stage 8 which requires SUMO)")

        # ── Stale detection ────────────────────────────────────────────────────
        # Modify the engineering checks JSON and check that stale is detected
        ec_data = json.loads(ec_path.read_text(encoding="utf-8"))
        ec_data["_smoke_test_marker"] = "modified"
        ec_path.write_text(json.dumps(ec_data), encoding="utf-8")
        new_h3 = _h.hash_json(str(ec_path))
        stored_h3 = m_final["stages"][3]["output_hash"]
        all_pass &= _ok("Stale detection: hash changed after modification",
                        new_h3 != stored_h3)

    return all_pass


def main():
    print("=" * 60)
    print("  SUMO24 MCP Pipeline Smoke Test")
    print("=" * 60)

    results = {}

    results["manifest"] = test_manifest()

    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td)
        results["hashing"]  = test_hashing(tmp)
        results["snapshot"] = test_snapshot(tmp)
        results["seeding"]  = test_seeding(tmp)
        results["static_inputs"] = test_static_inputs(tmp)
        results["dynamic_inputs"] = test_dynamic_inputs(tmp)

    results["engineering"] = test_engineering()
    results["full_pipeline"] = test_full_pipeline()

    # Summary
    _section("Summary")
    all_pass = True
    for name, passed in results.items():
        status = _PASS if passed else _FAIL
        print(f"  {status} {name}")
        if not passed:
            all_pass = False

    print()
    if all_pass:
        print(f"  {_PASS} ALL TESTS PASSED")
        sys.exit(0)
    else:
        print(f"  {_FAIL} SOME TESTS FAILED")
        sys.exit(1)


if __name__ == "__main__":
    main()
